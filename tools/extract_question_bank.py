"""
Extract the GEO audit question bank from the source Excel workbook into a
runtime config (config/question_bank.json).

Excel remains the single source of truth: edit the workbook tabs, re-run this
script, and the pipeline picks up the new questions / weights automatically.

For every scored category tab we capture, per question row:
    category          - the audit tab (== sheet name)
    subcategory       - grouping within the tab
    question          - the Yes/No question text
    weight            - the scoring weight (the "new weight" column the
                        workbook's own totals are computed from)
    how               - the original "How" hint (data source in the manual process)
    source            - classified pipeline data source (see SOURCE_RULES)
    mode              - "auto" (machine-scorable on free tiers) or "manual"
    reference_answer  - the human-filled Yes/No/N/A already in the workbook,
                        kept so we can measure auto-scoring accuracy
    check_id          - stable id
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

# The 11 scored category tabs (sheet name == audit category).
CATEGORY_TABS = [
    "Technical", "Site Speed", "Accessibility", "Schema", "Analytics",
    "Geotargetting", "Links", "Mobile", "Content", "GEO", "Architecture",
]

# Ordered (first match wins) classification rules.
# Each entry: (source, mode, [keywords matched against "How" + question text]).
SOURCE_RULES = [
    # --- genuinely manual / no free path -------------------------------
    ("manual_offpage", "manual",
     ["moz", "semrush", "backlink", "domain authority", "domain rating",
      "referring domain", "outreach"]),
    ("manual_listings", "manual",
     ["gmb", "google my business", "google business", "bing", "yellow pages",
      "citation", "local listing", "local review"]),
    ("manual_aio", "manual",
     ["perplexity", "gemini", "copilot", "ai overview", "aio", "chatgpt"]),
    ("manual_offpage", "manual",
     ["news website", "links from news", "links to the homepage from"]),
    # --- question-phrase overrides (beat the workbook's "Manual Check" default) ---
    # Detectable from crawled HTML / headers / robots / repeated requests:
    ("crawler", "auto",
     ["open graph", "twitter card", 'rel="me"', "sameas", "social media accounts",
      "social sharing", "breadcrumb", "flash", ".swf", "robots.txt",
      "javascript resources", "css resources", "trailing slash",
      "non-default protocol", "non-www", "capitalisation in the url", "cctld",
      "urls on the site absolute", "in the footer", "horizontal links",
      "blog on the same domain", "accordions or tabs", "hidden content",
      "white text", "infographics", "last updated", "click-to-call",
      "interstitial", "ip- or cookie-based", "howtostep",
      "name, address", "(nap)", "local landing pages include contact",
      "url structure of local landing", "ratio of links to deeper"]),
    # Judgeable by the LLM from extracted page text:
    ("llm", "auto",
     ["verifier", "attribution lines", "faq answers concise",
      "directly answer the headline", "contextual information included",
      "unique body copy", "target keywords excessively", "localised content",
      "relevant language for its target", "good and poor alt text",
      "brand and brand tagline", "avoid duplicate content across core"]),
    # --- free API collectors -------------------------------------------
    ("psi", "auto",
     ["pagespeed", "page insights", "page speed", "lighthouse", "crux",
      "chrome ux", "web vitals", "gtmetrix", "varvy", "keycdn", "http2",
      "mobile-friendly version", "mobile friendly version"]),
    ("gsc", "auto",
     ["search console", "gsc", "coverage report"]),
    ("ga4", "auto",
     ["google analytics", "ga4", "gtm", "tag management", "gachecker",
      "analytics"]),
    # --- LLM content judgment (from crawled page text) -----------------
    ("llm", "auto",
     ["readability", "hemingway", "flesch", "nlp", "fact-check", "fact check",
      "entity extraction", "tone audit", "tone/style", "manual content review",
      "manual tone", "manual review of", "manual product", "manual count",
      "manual scan", "customer chat", "people also ask", "ocr"]),
    # --- custom crawler (the Screaming Frog replacement) ---------------
    ("crawler", "auto",
     ["screaming frog", "crawl", "outlink", "orphan", "schema validator",
      "rich results", "structured data", "json-ld", "site command",
      "page source", "inspect element", "view source", "render", "sitemap",
      "site crawl", "image crawl", "anchor text export", "internal link"]),
    # --- everything else: subjective manual check ----------------------
    ("manual", "manual",
     ["manual", "copyscape"]),
]


def classify(how: str, question: str):
    blob = f"{how} {question}".lower()
    for source, mode, keywords in SOURCE_RULES:
        if any(k in blob for k in keywords):
            return source, mode
    return "manual", "manual"  # default: needs a human


def slug(text: str, n: int = 40) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s[:n]


def find_header_row(ws):
    """Return (row_index, {header_text_lower: col_index})."""
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [("" if c is None else str(c)).strip().lower() for c in row]
        if any("new question" in c or c == "question" for c in cells):
            return r, {c: i for i, c in enumerate(cells) if c}
    return None, {}


def col_for(headers, *names):
    for name in names:
        for text, idx in headers.items():
            if name in text:
                return idx
    return None


def to_weight(val):
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def main():
    wb_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("config/question_bank.json")
    wb = openpyxl.load_workbook(wb_path, data_only=True)

    bank = []
    for tab in CATEGORY_TABS:
        if tab not in wb.sheetnames:
            print(f"  ! tab not found, skipping: {tab}")
            continue
        ws = wb[tab]
        header_row, headers = find_header_row(ws)
        if not headers:
            print(f"  ! no header row in: {tab}")
            continue

        q_col = col_for(headers, "new question", "question")
        how_col = col_for(headers, "how")
        ans_col = col_for(headers, "yes/no", "yes / no")
        sub_col = col_for(headers, "subcategory")
        action_col = col_for(headers, "new action", "action")
        # Scoring weight = the column immediately after "New action"
        # (the workbook computes its category totals from this column).
        weight_col = (action_col + 1) if action_col is not None else None

        rows = list(ws.iter_rows(values_only=True))
        count = 0
        for row in rows[header_row:]:  # data rows after the header
            if q_col is None or q_col >= len(row):
                continue
            question = row[q_col]
            if not question or not str(question).strip():
                continue
            question = str(question).strip()
            how = str(row[how_col]).strip() if how_col is not None and row[how_col] else ""
            answer = str(row[ans_col]).strip() if ans_col is not None and row[ans_col] else ""
            subcat = str(row[sub_col]).strip() if sub_col is not None and row[sub_col] else ""
            weight = to_weight(row[weight_col]) if weight_col is not None and weight_col < len(row) else None
            if weight is None:
                continue  # rows without a weight aren't scored

            source, mode = classify(how, question)
            bank.append({
                "check_id": f"{slug(tab,12)}--{slug(question)}",
                "category": tab,
                "subcategory": subcat,
                "question": question,
                "weight": weight,
                "how": how,
                "source": source,
                "mode": mode,
                "reference_answer": answer,
            })
            count += 1
        print(f"  {tab:<16} {count:>3} questions")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(bank, indent=2, ensure_ascii=False), encoding="utf-8")

    # Summary
    print(f"\nTotal questions: {len(bank)}")
    by_source, auto_w, manual_w = {}, 0.0, 0.0
    for q in bank:
        by_source[q["source"]] = by_source.get(q["source"], 0) + 1
        if q["mode"] == "auto":
            auto_w += q["weight"]
        else:
            manual_w += q["weight"]
    print("By source:", json.dumps(by_source, indent=2))
    total_w = auto_w + manual_w
    print(f"Auto-scorable weight:   {auto_w:.0f} ({auto_w/total_w*100:.1f}%)")
    print(f"Manual-review weight:   {manual_w:.0f} ({manual_w/total_w*100:.1f}%)")
    print(f"Wrote -> {out_path}")


if __name__ == "__main__":
    main()
