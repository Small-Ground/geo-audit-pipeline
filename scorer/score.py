"""
Scorer — maps collected data to Yes/No/N·A per audit question, applies the
workbook's weighting math, and writes a populated Excel audit + score summary.

Two paths:
  * production: each category's evidence is sent to LiteLLM (OpenAI-compatible
    proxy, fronting Gemini) which returns Yes/No/N·A per question.
  * --offline:  deterministic rules (offline_rules.py) score the unambiguous
    crawler/PSI checks with no LLM/API key. Anything a rule can't decide is
    marked PENDING_LLM.

Scoring (reproduces the workbook exactly):
    category numerator   = Σ weight where answer == "Yes"
    category denominator = Σ weight where answer in {"Yes","No"}  (N·A dropped)
    category score       = numerator / denominator
    total weighting score= Σ numerators / Σ denominators
Manual-review and pending questions are excluded from the denominator and
reported separately as outstanding coverage.

Usage:
    python score.py --offline                 # no LLM, deterministic demo
    python score.py                           # LLM path (needs LiteLLM up)
Env (LLM path):
    LITELLM_BASE_URL (default http://localhost:4000/v1)
    LITELLM_API_KEY
    LLM_MODEL        (default gemini-2.0-flash)
"""

import argparse
import json
import os
import time
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import httpx
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from offline_rules import evaluate as offline_eval

ROOT = Path(__file__).resolve().parent.parent
MANUAL_MODES = {"manual", "manual_offpage", "manual_listings", "manual_aio"}


# --------------------------------------------------------------------------- #
# Evidence assembly
# --------------------------------------------------------------------------- #
def load(path, default):
    p = ROOT / path
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return default


def evidence_for(category, crawl, psi, gsc, ga4):
    ev = {"crawl_aggregates": crawl.get("aggregates", {}),
          "site_signals": crawl.get("site", {})}
    if category in ("Site Speed", "Mobile", "Content"):
        ev["pagespeed"] = {k: psi.get(k) for k in ("mobile", "desktop") if k in psi}
    if category in ("Accessibility", "Technical", "Geotargetting"):
        ev["search_console"] = gsc
    if category == "Analytics":
        ev["ga4"] = ga4
    if category in ("Content", "GEO", "Schema"):
        ev["page_samples"] = [
            {"url": p["url"], "title": p.get("title"), "h1": p.get("h1"),
             "word_count": p.get("word_count"), "schema": p.get("schema"),
             "text_sample": (p.get("text_sample") or "")[:700]}
            for p in crawl.get("page_sample", [])[:8]
        ]
    return ev


# --------------------------------------------------------------------------- #
# LLM path
# --------------------------------------------------------------------------- #
SYSTEM = (
    "You are a senior SEO/GEO technical auditor. For each audit question, decide "
    "strictly one of: Yes, No, N/A. 'Yes' ALWAYS means the site PASSES / meets "
    "best practice. Base every answer ONLY on the supplied evidence. If the "
    "evidence is insufficient to judge, answer N/A. Respond with JSON only: "
    '{"answers":[{"check_id":"...","answer":"Yes|No|N/A","confidence":0-1,'
    '"rationale":"<=20 words"}]}'
)


def ask_llm(category, evidence, questions):
    base = os.environ.get("LITELLM_BASE_URL", "http://localhost:4000/v1")
    key = os.environ.get("LITELLM_API_KEY", "sk-none")
    model = os.environ.get("LLM_MODEL", "gemini-2.0-flash")
    payload = {
        "model": model,
        "temperature": 0,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": SYSTEM},
            {"role": "user", "content": json.dumps({
                "category": category,
                "evidence": evidence,
                "questions": [{"check_id": q["check_id"], "question": q["question"]}
                              for q in questions],
            }, ensure_ascii=False)},
        ],
    }
    last_exc = None
    for attempt in range(5):
        r = httpx.post(f"{base}/chat/completions",
                       headers={"Authorization": f"Bearer {key}"},
                       json=payload, timeout=120)
        if r.status_code == 429:
            # respect Retry-After (Groq sends seconds), else exponential backoff
            wait = float(r.headers.get("retry-after", 0)) or min(2 ** attempt * 3, 30)
            print(f"    429 on {category} — waiting {wait:.0f}s (attempt {attempt + 1}/5)")
            time.sleep(wait + 0.5)
            continue
        try:
            r.raise_for_status()
        except httpx.HTTPStatusError as e:
            last_exc = e
            break
        content = r.json()["choices"][0]["message"]["content"]
        break
    else:
        raise httpx.HTTPError("exhausted retries (429)")
    if last_exc:
        raise last_exc
    content = content.strip().removeprefix("```json").removeprefix("```").removesuffix("```")
    answers = json.loads(content).get("answers", [])
    return {a["check_id"]: a for a in answers}


# --------------------------------------------------------------------------- #
# Scoring
# --------------------------------------------------------------------------- #
def score(bank, crawl, psi, gsc, ga4, offline):
    by_cat = defaultdict(list)
    for q in bank:
        by_cat[q["category"]].append(q)

    results = []
    for category, questions in by_cat.items():
        auto_qs = [q for q in questions if q["mode"] == "auto"]
        evidence = evidence_for(category, crawl, psi, gsc, ga4)

        llm_answers = {}
        if not offline and auto_qs:
            try:
                llm_answers = ask_llm(category, evidence, auto_qs)
            except (httpx.HTTPError, KeyError, json.JSONDecodeError) as e:
                print(f"  ! LLM failed for {category}: {e} -> PENDING_LLM")

        for q in questions:
            ans, conf, rat = "MANUAL_REVIEW", None, "Needs human review (no free data source)"
            if q["mode"] == "auto":
                if offline:
                    r = offline_eval(q["question"], crawl.get("aggregates", {}),
                                     crawl.get("site", {}), psi)
                    ans = r if r else "PENDING_LLM"
                    rat = "deterministic rule" if r else "no offline rule — needs LLM"
                    conf = 1.0 if r else None
                else:
                    a = llm_answers.get(q["check_id"])
                    if a:
                        ans, conf, rat = a.get("answer", "PENDING_LLM"), \
                            a.get("confidence"), a.get("rationale", "")
                    else:
                        ans, rat = "PENDING_LLM", "no LLM answer returned"
            results.append({**q, "answer": ans, "confidence": conf, "rationale": rat})
    return results


def summarise(results):
    cats = defaultdict(lambda: {"num": 0.0, "den": 0.0, "manual_w": 0.0,
                                "pending_w": 0.0, "answered": 0, "total": 0,
                                "ref_match": 0, "ref_total": 0})
    for r in results:
        c = cats[r["category"]]
        w, a = r["weight"], r["answer"]
        c["total"] += 1
        if a in ("Yes", "No"):
            c["den"] += w
            c["answered"] += 1
            if a == "Yes":
                c["num"] += w
            ref = (r.get("reference_answer") or "").strip().lower()
            if ref in ("yes", "no"):
                c["ref_total"] += 1
                if ref == a.lower():
                    c["ref_match"] += 1
        elif a == "MANUAL_REVIEW":
            c["manual_w"] += w
        elif a == "PENDING_LLM":
            c["pending_w"] += w
    for c in cats.values():
        c["score"] = (c["num"] / c["den"]) if c["den"] else None
    tot_num = sum(c["num"] for c in cats.values())
    tot_den = sum(c["den"] for c in cats.values())
    total = {"score": (tot_num / tot_den) if tot_den else None,
             "num": tot_num, "den": tot_den,
             "ref_match": sum(c["ref_match"] for c in cats.values()),
             "ref_total": sum(c["ref_total"] for c in cats.values())}
    return cats, total


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="3C3489")
FILLS = {"Yes": PatternFill("solid", fgColor="E4F3DE"),
         "No": PatternFill("solid", fgColor="FAECE7"),
         "N/A": PatternFill("solid", fgColor="F1EFE8"),
         "MANUAL_REVIEW": PatternFill("solid", fgColor="FAEEDA"),
         "PENDING_LLM": PatternFill("solid", fgColor="EEEDFE")}


def write_excel(results, cats, total, crawl, out_path):
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Score Summary"
    s.append(["GEO Technical Audit — automated (free-tier pipeline)"])
    s["A1"].font = Font(bold=True, size=14)
    s.append(["Domain", crawl.get("domain", "")])
    s.append(["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    s.append(["Pages crawled", crawl.get("pages_crawled", 0)])
    s.append([])
    header = ["Category", "Score", "Earned (Yes)", "Scored weight",
              "Auto-answered", "Manual pending", "LLM pending",
              "Total Qs", "vs human (match/total)"]
    s.append(header)
    hr = s.max_row
    for cell in s[hr]:
        cell.font, cell.fill = HEAD, HEAD_FILL
    for category in sorted(cats):
        c = cats[category]
        s.append([
            category,
            round(c["score"], 4) if c["score"] is not None else "n/a",
            round(c["num"]), round(c["den"]),
            f'{c["answered"]}/{c["total"]}',
            round(c["manual_w"]), round(c["pending_w"]), c["total"],
            f'{c["ref_match"]}/{c["ref_total"]}' if c["ref_total"] else "-",
        ])
    s.append([])
    s.append(["TOTAL WEIGHTING SCORE",
              round(total["score"], 4) if total["score"] is not None else "n/a",
              round(total["num"]), round(total["den"]), "", "", "", "",
              f'{total["ref_match"]}/{total["ref_total"]}' if total["ref_total"] else "-"])
    s[s.max_row][0].font = Font(bold=True)
    s[s.max_row][1].font = Font(bold=True)
    for col, w in zip("ABCDEFGHI", (26, 10, 12, 13, 14, 14, 12, 9, 20)):
        s.column_dimensions[col].width = w

    # one sheet per category
    by_cat = defaultdict(list)
    for r in results:
        by_cat[r["category"]].append(r)
    for category in sorted(by_cat):
        ws = wb.create_sheet(category[:31])
        cols = ["Subcategory", "Question", "Weight", "Source", "Answer",
                "Conf.", "Rationale", "Mode", "Human ref"]
        ws.append(cols)
        for cell in ws[1]:
            cell.font, cell.fill = HEAD, HEAD_FILL
        for r in by_cat[category]:
            ws.append([r["subcategory"], r["question"], r["weight"], r["source"],
                       r["answer"], round(r["confidence"], 2) if r["confidence"] else "",
                       r["rationale"], r["mode"], r.get("reference_answer", "")])
            ws[ws.max_row][4].fill = FILLS.get(r["answer"], PatternFill())
        for col, w in zip("ABCDEFGHI", (20, 70, 7, 14, 14, 6, 40, 14, 10)):
            ws.column_dimensions[col].width = w
        ws.column_dimensions["B"].width = 70
        for row in ws.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            row[6].alignment = Alignment(wrap_text=True, vertical="top")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--offline", action="store_true",
                    help="deterministic rules only, no LLM")
    ap.add_argument("--out", default="output/audit_populated.xlsx")
    args = ap.parse_args()

    bank = load("config/question_bank.json", [])
    crawl = load("output/crawl.json", {})
    psi = load("output/psi.json", {})
    gsc = load("output/gsc.json", {"available": False})
    ga4 = load("output/ga4.json", {"available": False})
    if not bank or not crawl:
        raise SystemExit("Missing config/question_bank.json or output/crawl.json")

    results = score(bank, crawl, psi, gsc, ga4, args.offline)
    cats, total = summarise(results)
    write_excel(results, cats, total, crawl, ROOT / args.out)

    print(f"\nTotal Weighting Score: "
          f"{round(total['score'], 4) if total['score'] is not None else 'n/a'} "
          f"({round(total['num'])}/{round(total['den'])})")
    if total["ref_total"]:
        print(f"Agreement with human answers: "
              f"{total['ref_match']}/{total['ref_total']} "
              f"({total['ref_match']/total['ref_total']*100:.0f}%)")
    for category in sorted(cats):
        c = cats[category]
        sc = round(c["score"], 3) if c["score"] is not None else "n/a"
        print(f"  {category:<16} score={sc!s:<6} answered={c['answered']}/{c['total']}"
              f"  manual_w={round(c['manual_w'])} pending_w={round(c['pending_w'])}")
    print(f"-> {args.out}")


if __name__ == "__main__":
    main()
